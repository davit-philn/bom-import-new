"""Quét hàng loạt file BOM (.xlsm) và gom số liệu để biết người dùng đang khai thế nào.

Chỉ ĐỌC file, không sửa. Kết quả: bom_scan_report.md (tóm tắt) + bom_scan_report.json (chi tiết từng file).
Chạy:  python tools/scan_boms.py "D:\\boho\\BOM"
"""
import json
import re
import sys
import warnings
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

IGNORE_NOTE = {"", None}


def txt(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def scan_file(path):
    info = {"file": path.name, "folder": path.parent.name, "errors": []}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wv = openpyxl.load_workbook(path, data_only=True)
        wf = openpyxl.load_workbook(path)
    info["sheets"] = [(ws.title, ws.sheet_state) for ws in wf.worksheets]
    if "BOM_Phan II" not in wv.sheetnames:
        info["errors"].append("không có sheet BOM_Phan II")
        return info
    v2, f2 = wv["BOM_Phan II"], wf["BOM_Phan II"]
    info["header"] = {k: txt(v2[c].value) for k, c in
                      dict(project="E4", order="P4", code="E6", muc="E7", structure="P7", slsp="Y3", doorType="Q7").items()}
    info["dims"] = f2.dimensions
    # phiên bản công thức chỉ cạnh / trọng lượng
    s13 = str(f2["S13"].value or "")
    info["edge_formula"] = ("+0.04/cạnh" if "0.04" in s13 else "+80" if "+80" in s13 else "+40" if "+40" in s13 else "chuẩn")
    y13 = str(f2["Y13"].value or "")
    info["weight_formula"] = re.sub(r"[A-Z]+\d+", "#", y13)[:60]
    info["has_pivot_Q"] = any(f2.cell(r, 17).value for r in range(12, 16))

    # dòng dữ liệu mục A
    start = next((r for r in range(1, 60) if txt(v2.cell(r, 4).value) == "VẬT TƯ GỖ, VẬT TƯ TẤM"), None)
    rows = []
    if start:
        for r in range(start + 1, v2.max_row + 1):
            d = txt(v2.cell(r, 4).value)
            if d and d.upper() in ("VẬT TƯ FOAM SOFA", "TỔNG CỘNG:"):
                break
            if d and any(v2.cell(r, c).value not in (None, "") for c in (5, 6, 7)):
                rows.append(r)
    info["partsA"] = len(rows)
    notes, xs, kfmt, mfmt, stt_kind = Counter(), Counter(), Counter(), Counter(), Counter()
    parents = 0
    for r in rows:
        o = txt(v2.cell(r, 15).value)
        notes[(o or "")[:60]] += 1
        xv, xf = v2.cell(r, 24).value, f2.cell(r, 24).value
        xs[("công thức" if isinstance(xf, str) and str(xf).startswith("=") else "gõ tay" if xf is not None else "trống") + ":" + str(xv)] += 1
        kfmt[re.sub(r"\d", "#", txt(v2.cell(r, 11).value) or "")] += 1
        mfmt[re.sub(r"\d", "#", txt(v2.cell(r, 13).value) or "")] += 1
        s = txt(v2.cell(r, 3).value) or ""
        stt_kind["nguyên" if re.fullmatch(r"\d+", s) else "thập phân" if re.fullmatch(r"\d+\.\d+", s) else "khác/trống"] += 1
        if xv in (1, 2, "1", "2"):
            parents += 1
    info["notes"] = notes.most_common(15)
    info["pldg"] = xs.most_common(10)
    info["edges_fmt"] = kfmt.most_common(6)
    info["faces_fmt"] = mfmt.most_common(6)
    info["stt_kind"] = dict(stt_kind)
    info["rows_pldg_1_2"] = parents
    info["finish_codes"] = [txt(v2.cell(r, 22).value) for r in range(2, 9) if txt(v2.cell(r, 22).value)]
    info["units"] = Counter(txt(v2.cell(r, 17).value) for r in rows).most_common(5)

    # Phần III
    if "BOM_Phan III" in wv.sheetnames:
        v3 = wv["BOM_Phan III"]
        sec = [(r, txt(v3.cell(r, 4).value)) for r in range(10, v3.max_row + 1)
               if txt(v3.cell(r, 4).value) and txt(v3.cell(r, 4).value).isupper() and not txt(v3.cell(r, 5).value)
               and len(txt(v3.cell(r, 4).value)) > 6 and not txt(v3.cell(r, 4).value).startswith("NVL")]
        info["p3_sections"] = [n for _, n in sec]
    # Phần IV
    if "BOM_Phan IV" in wv.sheetnames:
        v4 = wv["BOM_Phan IV"]
        info["packaging_standard"] = txt(v4["O13"].value) or txt(v4["O12"].value)
        info["coef_S6"] = v4["S6"].value
        kien = [r for r in range(13, min(v4.max_row, 1500)) if txt(v4.cell(r, 3).value) and str(v4.cell(r, 3).value).upper().startswith("KIỆN")]
        info["p4_kien_count"] = len(kien)
    # lỗi công thức
    err = Counter()
    for ws in wf.worksheets:
        if ws.sheet_state != "visible":
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "#REF!" in c.value:
                    err[ws.title] += 1
    info["ref_errors"] = dict(err)
    return info


def main(folder):
    files = sorted(p for p in Path(folder).rglob("*.xls*") if not p.name.startswith("~$"))
    print("Tìm thấy", len(files), "file")
    out = []
    for i, p in enumerate(files, 1):
        try:
            out.append(scan_file(p))
            print("[%d/%d] %s" % (i, len(files), p.name))
        except Exception as e:  # noqa: BLE001
            out.append({"file": p.name, "folder": p.parent.name, "errors": [repr(e)]})
            print("[%d/%d] LỖI %s: %s" % (i, len(files), p.name, e))
    Path("bom_scan_report.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

    agg = defaultdict(Counter)
    for o in out:
        if "header" not in o:
            continue
        agg["Công thức chỉ cạnh"][o["edge_formula"]] += 1
        agg["Công thức trọng lượng (cột Y)"][o["weight_formula"]] += 1
        agg["Tiêu chuẩn bao bì (O13)"][str(o.get("packaging_standard"))] += 1
        agg["Hệ số tỉ trọng (S6)"][str(o.get("coef_S6"))] += 1
        agg["Mục Phần III"][" | ".join(o.get("p3_sections", []))] += 1
        agg["Kết cấu (P7)"][str(o["header"]["structure"])] += 1
        agg["Loại cửa (Q7)"][str(o["header"]["doorType"])] += 1
        agg["Có lỗi #REF!"][str(bool(o["ref_errors"]))] += 1
        for k, n in o["notes"]:
            agg["Ghi chú/giao hàng (cột O)"][k] += n
        for k, n in o["pldg"]:
            agg["PLDG (cột X)"][k] += n
        for k, n in o["edges_fmt"]:
            agg["Định dạng số cạnh (K)"][k] += n
        for k, n in o["faces_fmt"]:
            agg["Định dạng mặt sơn (M)"][k] += n
        for k, n in o["stt_kind"].items():
            agg["Kiểu STT (cột C)"][k] += n
    lines = ["# Báo cáo quét BOM", "", "Số file: %d (lỗi đọc: %d)" % (len(out), sum(1 for o in out if "header" not in o)), ""]
    for title, cnt in agg.items():
        lines += ["## " + title, ""] + ["- `%s`: %d" % (k, n) for k, n in cnt.most_common(25)] + [""]
    lines += ["## Từng file (số chi tiết A / số dòng PLDG 1-2 / số kiện Phần IV)", ""]
    for o in out:
        if "header" in o:
            lines.append("- %s / %s: A=%d, PLDG1-2=%d, kiện=%s, chỉ cạnh=%s" % (o["folder"], o["file"], o["partsA"], o["rows_pldg_1_2"], o.get("p4_kien_count"), o["edge_formula"]))
    Path("bom_scan_report.md").write_text("\n".join(lines), encoding="utf-8")
    print("\nXong. Gửi file bom_scan_report.md (và .json nếu muốn) cho Claude.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit('Cách dùng: python tools/scan_boms.py "đường-dẫn-folder-BOM"')
    main(sys.argv[1])
