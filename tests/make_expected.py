"""Tạo fixture đối chiếu: đầu vào + giá trị Excel đã tính (cache) của Phần II mục A trong 3 file mẫu."""
import glob, json, sys, warnings
from pathlib import Path
import openpyxl
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import importer

COL = {"unit": 17, "mass": 18, "edgeLen": 19, "areaFinish": 20, "areaPrime": 21, "weight": 25, "pldg": 24}
for f in sorted(glob.glob(str(Path(__file__).parent / "samples" / "*.xlsm"))):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(f, data_only=True)["BOM_Phan II"]
    start = next(r for r in range(1, 40) if ws.cell(r, 4).value == "VẬT TƯ GỖ, VẬT TƯ TẤM")
    end = next(r for r in range(start + 1, ws.max_row) if ws.cell(r, 4).value in ("VẬT TƯ FOAM SOFA",) or ws.cell(r, 3).value == "B")
    cases = []
    for r in range(start + 1, end):
        if not ws.cell(r, 4).value:
            continue
        row = {k: ws.cell(r, c).value for k, c in importer.PART_COLS.items()}
        exp = {k: ws.cell(r, c).value for k, c in COL.items()}
        cases.append({"row": r, "input": row, "expected": exp})
    out = Path(__file__).parent / "fixtures" / ("expected_" + Path(f).name[:8] + ".json")
    out.write_text(json.dumps(cases, ensure_ascii=False, indent=1), encoding="utf-8")
    print(out.name, len(cases))
