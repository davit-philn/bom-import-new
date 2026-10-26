"""Đọc file BOM .xlsm hiện có (cả 3 phiên bản mẫu) thành state JSON của web."""
import io
import re
import warnings

import openpyxl

PART_COLS = {"stt": 3, "name": 4, "thick": 5, "width": 6, "length": 7, "qty": 8, "material": 9,
             "edgeCode": 10, "edges": 11, "finishCode": 12, "faces": 13, "primeFaces": 14,
             "delivery": 15}
NUMERIC = {"thick", "width", "length", "qty", "qtyFactory", "qtySite", "qtySub"}
SUPPLY_COLS = {"code": 4, "name": 5, "unit": 6, "spec": 7, "material": 8, "supplier": 9,
               "qtyFactory": 10, "qtySite": 11, "qtySub": 12, "note": 13, "image": 14}
SUPPLY_EXTRA = {"thick": 17, "width": 18, "length": 19, "qty": 20, "pldg": 21}
PART_SECTIONS = [("partsA", "VẬT TƯ GỖ, VẬT TƯ TẤM"), ("partsB", "VẬT TƯ FOAM SOFA"),
                 ("partsC", "VẬT TƯ VẢI, DA, SIMILI")]
SUPPLY_SECTIONS = [("metal", "VẬT TƯ KIM LOẠI"), ("metalWork", "GIA CÔNG KIM LOẠI"),
                   ("aux", "VẬT TƯ PHỤ"), ("accessories", "PHỤ KIỆN"),
                   ("subcon", "THẦU PHỤ GIA CÔNG")]


def _clean(v, numeric=False):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if numeric:
        return int(v) if float(v) == int(v) else round(float(v), 6)
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v


def _stt(v):
    if v is None:
        return None
    if isinstance(v, float):
        s = ("%.10g" % v)
        return s
    return str(v).strip() or None


def _find(ws, col, title, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip().upper() == title:
            return r
    return None


def _text(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def import_xlsm(raw):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    for n in ("BOM_Phan II", "BOM_Phan III"):
        if n not in wb.sheetnames:
            raise ValueError("Không phải file BOM đúng mẫu: thiếu sheet %s" % n)
    ws2, ws3 = wb["BOM_Phan II"], wb["BOM_Phan III"]

    codes = [_text(ws2.cell(r, 22).value) for r in range(2, 9)]
    codes = [c for c in codes if c]
    if not codes and _text(ws2["P8"].value):
        codes = [c.strip() for c in str(ws2["P8"].value).split(",") if c.strip()]
    header = {
        "project": _text(ws2["E4"].value), "order": _text(ws2["P4"].value),
        "name": _text(ws2["E5"].value), "size": _text(ws2["P5"].value),
        "code": _text(ws2["E6"].value), "qty": _text(ws2["P6"].value),
        "muc": _text(ws2["E7"].value), "structure": _text(ws2["P7"].value),
        "material": _text(ws2["E8"].value), "slsp": _clean(ws2["Y3"].value, True),
        "doorType": _text(ws2["Q7"].value), "finishCodes": codes,
    }
    state = {"header": header}

    # Phần II
    heads = []
    for key, title in PART_SECTIONS:
        r = _find(ws2, 4, title)
        if r:
            heads.append((r, key))
    heads.sort()
    total = _find(ws2, 4, "TỔNG CỘNG:", start=heads[-1][0] if heads else 1) or ws2.max_row + 1
    for key, _ in PART_SECTIONS:
        state[key] = []
    ends = [h[0] for h in heads[1:]] + [total]
    for (start, key), end in zip(heads, ends):
        for r in range(start + 1, end):
            row = {}
            for f, c in PART_COLS.items():
                v = ws2.cell(r, c).value
                row[f] = _stt(v) if f == "stt" else _clean(v, f in NUMERIC)
            if not any(row.get(k) for k in ("name", "thick", "width", "length")):
                continue
            x = _clean(ws2.cell(r, 24).value, True)
            note = (row.get("delivery") or "").strip().upper()
            derived = None if not note else (1 if note == "HOÀN THIỆN ĐI MODULE" else 2)
            if x in (1, 2) and x != derived:
                row["pldg"] = x  # người dùng gõ tay PLDG khác công thức mặc định
            if key != "partsA":
                unit = _text(ws2.cell(r, 17).value)
                if unit:
                    row["unit"] = unit
            state[key].append({k: v for k, v in row.items() if v is not None})

    # chữ ký cuối Phần II
    sign = {}
    k = None
    for r in range(total, ws2.max_row + 1):
        v = ws2.cell(r, 7).value
        if isinstance(v, str) and "duyệt" in v.lower():
            k = r
        if isinstance(ws2.cell(r, 16).value, str) and ws2.cell(r, 16).value.strip().startswith("Ngày"):
            sign["dateText"] = ws2.cell(r, 16).value.strip()
    if k:
        sign["approver"] = _text(ws2.cell(k + 4, 7).value)
        sign["author"] = _text(ws2.cell(k + 4, 16).value)
    state["sign"] = {a: b for a, b in sign.items() if b}

    # Phần III
    sup_heads = []
    for key, title in SUPPLY_SECTIONS:
        r = _find(ws3, 4, title)
        if r:
            sup_heads.append((r, key))
    sup_heads.sort()
    state["supplies"] = {k: [] for k, _ in SUPPLY_SECTIONS}
    for i, (start, key) in enumerate(sup_heads):
        end = sup_heads[i + 1][0] if i + 1 < len(sup_heads) else None
        r = start + 1
        blank_run = 0
        while r < ws3.max_row + 1 and (end is None or r < end):
            row = {}
            cols = dict(SUPPLY_COLS)
            if key == "metalWork":
                cols.update(SUPPLY_EXTRA)
            for f, c in cols.items():
                v = ws3.cell(r, c).value
                row[f] = _clean(v, f in NUMERIC)
            tail = str(ws3.cell(r, 4).value or "").strip().lower()
            if end is None and (tail.startswith("kiểm duyệt") or tail.startswith("ngày")):
                break
            if any(row.get(k) for k in ("code", "name")):
                state["supplies"][key].append({a: b for a, b in row.items() if b is not None})
                blank_run = 0
            else:
                blank_run += 1
                if end is None and blank_run > 40:
                    break
            r += 1

    # Phần I + đóng gói
    if "BOM_Phan I" in wb.sheetnames:
        w1 = wb["BOM_Phan I"]
        hist = []
        for r in range(11, 17):
            row = {"no": _stt(w1.cell(r, 3).value), "basis": _text(w1.cell(r, 4).value),
                   "content": _text(w1.cell(r, 5).value),
                   "date": _text(w1.cell(r, 7).value), "by": _text(w1.cell(r, 8).value)}
            if row["basis"] or row["content"]:
                hist.append({a: b for a, b in row.items() if b})
        state["phieu"] = {"history": hist}
        if not state["sign"].get("approver"):
            state["sign"]["approver"] = _text(w1["D24"].value)
        if not state["sign"].get("author"):
            state["sign"]["author"] = _text(w1["G24"].value)
        if not state["sign"].get("dateText"):
            state["sign"]["dateText"] = _text(w1["G19"].value)
    if "BOM_Phan IV" in wb.sheetnames:
        std = _text(wb["BOM_Phan IV"]["O13"].value)
        if std:
            state["packaging"] = {"standard": std}
    return state
